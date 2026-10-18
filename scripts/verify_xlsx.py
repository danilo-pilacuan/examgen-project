"""Verification of the exported Excel (runs inside the container)."""
import sys
sys.path.insert(0, "/app")
from openpyxl import load_workbook

wb = load_workbook("/app/output/preguntas_20260917_071733.xlsx")
ws = wb.active
print("Rows:", ws.max_row, "Columns:", ws.max_column)

# Find a row with options (multiple_choice)
found = False
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[6] == "multiple_choice":
        print("Question:", str(row[7])[:60])
        print("Options:", row[8])
        print("Answer:", row[9])
        print("Human verif.:", row[12])
        found = True
        break
if not found:
    print("No rows with multiple_choice")

# Count question types
from collections import Counter
types = Counter(row[6] for row in ws.iter_rows(min_row=2, values_only=True))
print("Types:", dict(types))
